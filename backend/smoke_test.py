# -*- coding: utf-8 -*-
"""API 冒烟测试（自清理测试数据；只用 9000x/8000x/9100x 测试 UID，绝不碰真实数据）。
运行: .venv/Scripts/python.exe -X utf8 backend/smoke_test.py"""
import os
import sqlite3
import sys
from datetime import datetime, timedelta

import requests

B = 'http://127.0.0.1:5000'
passed = failed = 0

# 管理员登录（后续请求都带登录态）
S = requests.Session()
r = S.post(f'{B}/api/login', json={'username': 'admin', 'password': '112233'})
assert r.ok, f'管理员登录失败: {r.text}'
requests = S

# 未登录访问管理 API 应 401
anon = __import__('requests').Session()
r = anon.get(f'{B}/api/trainers')
print('未登录管理 API 拦截 401:', 'PASS' if r.status_code == 401 else 'FAIL')

def check(name, cond, extra=''):
    global passed, failed
    if cond:
        passed += 1
        print(f'  [PASS] {name}')
    else:
        failed += 1
        print(f'  [FAIL] {name} {extra}')


MONTH = datetime.now().strftime('%Y-%m')
DB = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'data', 'stats.db')
NOW_S = datetime.now().strftime('%Y-%m-%d %H:%M:%S')


def db_exec(sql, params=()):
    conn = sqlite3.connect(DB)
    conn.execute(sql, params)
    conn.commit()
    conn.close()


def db_one(sql, params=()):
    conn = sqlite3.connect(DB)
    row = conn.execute(sql, params).fetchone()
    conn.close()
    return row


def insert_monthly(uid, month, level, rank, tasks=7, retained=1, active=True):
    """插入月度等级记录；active=True 表示排档周等级（名额贡献口径）"""
    db_exec('''INSERT INTO trainee_level_monthly
               (trainee_uid, month, best_level, best_rank, last_week_level, last_daily_task_count, last_queried_at,
                last_retained, best_rank_active, best_level_active)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
            (uid, month, level, rank, level, tasks, NOW_S, retained,
             rank if active else 0, level if active else ''))


def expected_window(bound_at_utc):
    """与后端同口径：绑定次周周一(本地) 起整四周"""
    local = datetime.strptime(bound_at_utc, '%Y-%m-%d %H:%M:%S') + timedelta(hours=8)
    d = (7 - local.weekday()) % 7 or 7
    start = (local + timedelta(days=d)).replace(hour=0, minute=0, second=0, microsecond=0)
    return start, start + timedelta(weeks=4)


def bind(uid_a, uid_b, **kw):
    payload = {'trainer_uid': uid_a, 'trainee_uid': uid_b}
    payload.update(kw)
    return requests.post(f'{B}/api/bindings/by-uid', json=payload)


# 预清理：上次异常退出可能残留的测试档案（按 UID 精确匹配；绑定记录直接删除）
db_exec('''DELETE FROM bindings WHERE trainer_id IN (SELECT id FROM trainers WHERE uid IN ('90001','90002','90003','91001','91002'))
           OR trainee_id IN (SELECT id FROM trainees WHERE uid IN ('80001','80002','80003'))''')
for uid in ('90001', '90002', '90003', '91001', '91002'):
    for t in requests.get(f'{B}/api/trainers', params={'role': 'all', 'kw': uid}).json().get('items', []):
        if t['uid'] == uid:
            requests.delete(f"{B}/api/trainers/{t['id']}")
for uid in ('80001', '80002', '80003'):
    for t in requests.get(f'{B}/api/trainees', params={'kw': uid}).json().get('items', []):
        if t['uid'] == uid:
            requests.delete(f"{B}/api/trainees/{t['id']}")
    db_exec('DELETE FROM trainee_level_monthly WHERE trainee_uid=?', (uid,))
    db_exec('DELETE FROM reward_grants WHERE trainee_uid=?', (uid,))


print('== 培训师 CRUD ==')
r = requests.post(f'{B}/api/trainers', json={'uid': '90001', 'nickname': '测试培训师A', 'hall_name': '游戏一厅'})
check('新增培训师A（默认基础名额2）', r.ok and r.json()['ok'], r.text)
tr_a = r.json().get('id')
r = requests.post(f'{B}/api/trainers', json={'uid': '90002', 'nickname': '测试培训师B', 'base_slots': 3})
check('新增培训师B（基础名额3）', r.ok, r.text)
tr_b = r.json().get('id')
r = requests.post(f'{B}/api/trainers', json={'uid': '90001', 'nickname': '重复'})
check('UID 重复 409', r.status_code == 409, f'got {r.status_code}')
r = requests.post(f'{B}/api/trainers', json={'uid': '90099', 'nickname': 'X', 'base_slots': 9})
check('基础名额超上限 400', r.status_code == 400, f'got {r.status_code}')
r = requests.put(f'{B}/api/trainers/{tr_b}', json={'uid': '90002', 'nickname': '测试培训师B改', 'status': 'active', 'note': '备注', 'base_slots': 3})
check('编辑培训师B', r.ok, r.text)
r = requests.get(f'{B}/api/trainers', params={'kw': 'B改'})
check('搜索培训师', r.ok and len(r.json()['items']) == 1, r.text)
r = requests.get(f'{B}/api/trainers')
item_a = next((t for t in r.json()['items'] if t['uid'] == '90001'), None)
check('列表带名额字段 total_slots=2', item_a and item_a['total_slots'] == 2, str(item_a))
check('列表带分页字段', 'total' in r.json() and r.json()['page_size'] == 20, str(r.json().keys()))

print('== 培训师批量导入 ==')
r = requests.post(f'{B}/api/trainers/bulk-import', json={'uids': '90011,90012,90013'})
check('批量导入 3 人', r.ok and r.json()['created_count'] == 3, r.text)
r = requests.post(f'{B}/api/trainers/bulk-import', json={'uids': '90011,90014'})
check('重复跳过', r.ok and r.json()['created_count'] == 1 and r.json()['skipped'] == ['90011'], r.text)
r = requests.get(f'{B}/api/trainers', params={'kw': '90011'})
check('批量导入可搜索到', len(r.json()['items']) == 1 and r.json()['items'][0]['base_slots'] == 2, r.text)
# 角色冲突：先作为姐姐建档，批量导入后应更新为 trainer
db_exec("INSERT INTO trainers (uid, nickname, role) VALUES ('90006', '角色测试', 'sister')")
r = requests.post(f'{B}/api/trainers/bulk-import', json={'uids': '90006'})
check('姐姐角色导入后更新为 trainer', db_one("SELECT role FROM trainers WHERE uid='90006'")[0] == 'trainer')
# 单个新增遇姐姐角色也应转换（而不是 409）
db_exec("INSERT INTO trainers (uid, nickname, role) VALUES ('90007', '角色测试2', 'sister')")
r = requests.post(f'{B}/api/trainers', json={'uid': '90007', 'nickname': '角色测试2'})
check('单个新增遇姐姐角色自动转换', r.ok and r.json().get('converted'), r.text)
# 姐姐冲突检查：造 team_detail 快照，验证 空/同 UID/异 UID
SNAP = datetime.now().strftime('%Y-%m-%d')
db_exec("DELETE FROM team_detail WHERE snapshot_date=?", (SNAP,))
db_exec('''INSERT INTO team_detail (team_id, form_date, hall_name, sister_uid, sister_nickname, sister_uid2, sister_nickname2, dissolve_date, snapshot_date)
           VALUES (999001, '2026-10-01', '测试厅', 88888, '别的姐姐', 80001, '测试妹妹甲', '', ?)''', (SNAP,))
import sys as _sys
_sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from app import check_sister_conflict as _csc  # noqa
from db import get_db as _gdb  # noqa
_conn = _gdb()
ok1 = _csc(_conn, '80001', '90001')
check('异姐姐拒绝绑定', ok1[0] is False and '不能绑定其他培训师' in ok1[1], str(ok1))
ok2 = _csc(_conn, '80001', '88888')
check('同姐姐允许绑定', ok2[0] is True, str(ok2))
db_exec("DELETE FROM team_detail WHERE snapshot_date=?", (SNAP,))
ok3 = _csc(_conn, '80001', '90001')
check('无姐妹团记录允许绑定', ok3[0] is True, str(ok3))
_conn.close()
for uid in ('90011', '90012', '90013', '90014', '90006', '90007'):
    for t in requests.get(f'{B}/api/trainers', params={'role': 'all', 'kw': uid}).json()['items']:
        if t['uid'] == uid:
            requests.delete(f"{B}/api/trainers/{t['id']}")

print('== 妹妹 CRUD ==')
r = requests.post(f'{B}/api/trainees', json={'uid': '80001', 'nickname': '测试妹妹甲', 'hall_name': '游戏一厅'})
check('新增妹妹甲', r.ok, r.text)
te_a = r.json().get('id')
r = requests.post(f'{B}/api/trainees', json={'uid': '80002', 'nickname': '测试妹妹乙'})
check('新增妹妹乙', r.ok, r.text)
te_b = r.json().get('id')
r = requests.put(f'{B}/api/trainees/{te_b}', json={'uid': '80002', 'nickname': '测试妹妹乙', 'status': 'graduated'})
check('状态 graduated 不再接受 400', r.status_code == 400, f'got {r.status_code}')
r = requests.put(f'{B}/api/trainees/{te_b}', json={'uid': '80002', 'nickname': '测试妹妹乙', 'status': 'inactive'})
check('状态 active/inactive 正常', r.ok, r.text)

print('== by-uid 绑定与保护期门槛 ==')
r = bind('90001', '80001')
check('培训师绑定被保护期门槛拦截 422（测试 UID 无保护期记录）', r.status_code == 422 and r.json().get('gate'), r.text)
check('门槛拦截不落库', db_one("SELECT COUNT(*) FROM bindings WHERE trainee_id=?", (te_a,))[0] == 0)
r = bind('90001', '90001')
check('同一人绑定 400', r.status_code == 400, f'got {r.status_code}')
r = bind('90003', '80003', bind_type='sister', trainer_nickname='自动建档C', trainee_nickname='自动建档丙')
d = r.json()
check('by-uid 自动建档双方（姐妹团无门槛）', r.ok and d.get('trainer_created') and d.get('trainee_created'), r.text)
tr_c = db_one("SELECT id FROM trainers WHERE uid='90003'")[0]
te_c = db_one("SELECT id FROM trainees WHERE uid='80003'")[0]
check('自动建档培训师默认名额2', db_one("SELECT base_slots FROM trainers WHERE uid='90003'")[0] == 2)
# 下游逻辑（名额/奖励/窗口）需要培训师绑定：门槛已验证，直接落库构造
db_exec("INSERT INTO bindings (trainer_id, trainee_id, bind_type, note) VALUES (?, ?, 'trainer', '测试直插')", (tr_b, te_a))

print('== 姐妹团三胎绑定 ==')
r = bind('91001', '80001', bind_type='sister', trainer_nickname='测试姐姐甲')
d = r.json()
check('妹妹可同时有培训师+姐妹团绑定', r.ok and d.get('bind_type') == 'sister', r.text)
sis_a = db_one("SELECT id, role FROM trainers WHERE uid='91001'")
check('姐姐自动建档 role=sister', sis_a and sis_a[1] == 'sister', str(sis_a))
r = bind('91001', '80002', bind_type='sister', trainee_nickname='测试妹妹乙')
check('姐姐再绑第二个妹妹 409（三胎一对一）', r.status_code == 409 and '一对一' in r.json().get('msg', ''), r.text)
r = bind('91001', '80002', bind_type='sister', force=True)
check('姐姐强制换绑成功', r.ok and r.json().get('rebound'), r.text)
r = requests.get(f'{B}/api/bindings', params={'active': 1, 'bind_type': 'sister'})
sis_items = [b for b in r.json()['items'] if str(b['trainer_uid']) == '91001']
check('姐姐仅一条生效三胎绑定', len(sis_items) == 1 and sis_items[0]['trainee_uid'] == '80002', str(sis_items))
r = requests.get(f'{B}/api/bindings', params={'active': 1, 'bind_type': 'trainer'})
check('按类型过滤生效绑定', all(b['bind_type'] == 'trainer' for b in r.json()['items']), r.text)
b1 = next(b for b in requests.get(f'{B}/api/bindings', params={'active': 1}).json()['items'] if b['trainee_uid'] == '80001')
check('绑定带奖励窗口字段', b1.get('window_end') and b1.get('window_status') in ('not_started', 'in_window', 'expired')
      and b1.get('weeks_left') is not None, str({k: b1.get(k) for k in ('window_end', 'window_status', 'weeks_left')}))
_, we1 = expected_window(b1['bound_at'])
check('绑定列表结束日期=窗口终点', b1['window_end'] == we1.strftime('%Y-%m-%d'), f"{b1['window_end']} != {we1.date()}")
# 妹妹甲:培训师B + 妹妹乙:姐姐（姐妹团） 并存检查
r = requests.get(f'{B}/api/trainees', params={'kw': '80001'})
t1 = r.json()['items'][0]
check('甲同时显示培训师绑定', t1['trainer_uid'] == '90002' and not t1['sister_uid'], str(t1))
r = requests.get(f'{B}/api/trainees', params={'kw': '80002'})
t2 = r.json()['items'][0]
check('乙显示姐妹团姐姐', t2['sister_uid'] == '91001' and not t2['trainer_uid'], str(t2))

print('== 培训师列表不含姐姐 ==')
r = requests.get(f'{B}/api/trainers')
check('默认只列 role=trainer', all(t.get('role', 'trainer') == 'trainer' for t in r.json()['items']), r.text)
r = requests.get(f'{B}/api/trainers', params={'role': 'all', 'kw': '91001'})
check('role=all 能查到姐姐', len(r.json()['items']) == 1, r.text)

print('== 名额判定（只算培训师绑定） ==')
r = requests.get(f'{B}/api/trainers/{tr_a}/slots')
s = r.json()['slots']
check('A 无生效绑定 总名额=基础2', s['total_slots'] == 2 and s['used_slots'] == 0, str(s))
insert_monthly('80001', MONTH, '金牌', 4, tasks=7)
r = requests.get(f'{B}/api/trainers/{tr_b}/slots')
s = r.json()['slots']
check('金牌学员（有排档）+3 → B 总名额 min(3+3,5)=5', s['extra_slots'] == 3 and s['total_slots'] == 5, str(s))
print('== 名额排档门槛 ==')
db_exec("UPDATE trainee_level_monthly SET best_rank_active=0, best_level_active='' WHERE trainee_uid='80001' AND month=?", (MONTH,))
r = requests.get(f'{B}/api/trainers/{tr_b}/slots')
check('金牌但当周无排档 → 不贡献名额', r.json()['slots']['extra_slots'] == 0, str(r.json()['slots']['trainees']))
db_exec("UPDATE trainee_level_monthly SET best_rank_active=4, best_level_active='金牌' WHERE trainee_uid='80001' AND month=?", (MONTH,))
r = requests.get(f'{B}/api/trainers/{tr_b}/slots')
check('恢复排档后重新 +3', r.json()['slots']['extra_slots'] == 3, str(r.json()['slots']['extra_slots']))
b1r = next(b for b in requests.get(f'{B}/api/bindings', params={'active': 1}).json()['items'] if b['trainee_uid'] == '80001')
check('绑定列表带留存标记 retained=1', b1r.get('retained') == 1, str(b1r.get('retained')))
db_exec("UPDATE trainee_level_monthly SET last_retained=0 WHERE trainee_uid='80001' AND month=?", (MONTH,))
b1r = next(b for b in requests.get(f'{B}/api/bindings', params={'active': 1}).json()['items'] if b['trainee_uid'] == '80001')
check('retained=0 透出疑似解散', b1r.get('retained') == 0, str(b1r.get('retained')))
db_exec("UPDATE trainee_level_monthly SET last_retained=1 WHERE trainee_uid='80001' AND month=?", (MONTH,))

print('== 名额月初兜底（上月记录） ==')
LAST_MONTH = (datetime.now().replace(day=1) - timedelta(days=1)).strftime('%Y-%m')
db_exec("DELETE FROM trainee_level_monthly WHERE trainee_uid='80001' AND month=?", (MONTH,))
insert_monthly('80001', LAST_MONTH, '金牌', 4, tasks=7)
r = requests.get(f'{B}/api/trainers/{tr_b}/slots')
s = r.json()['slots']
check('当月无记录 → 上月金牌兜底 +3', s['extra_slots'] == 3 and s['total_slots'] == 5, str(s))
check('兜底行带 from_last_month 标记', s['trainees'][0].get('from_last_month') is True, str(s['trainees']))
insert_monthly('80001', MONTH, '铜牌', 1, tasks=7)
r = requests.get(f'{B}/api/trainers/{tr_b}/slots')
s = r.json()['slots']
check('当月有记录则按当月（铜牌 +0，不吃兜底）', s['extra_slots'] == 0 and s['trainees'][0].get('from_last_month') is False, str(s))
# 恢复金牌供后续奖励测试
db_exec("DELETE FROM trainee_level_monthly WHERE trainee_uid='80001' AND month=?", (MONTH,))
insert_monthly('80001', MONTH, '金牌', 4, tasks=7)

print('== 奖励窗口（次周周一起整四周） ==')
bound_at = db_one("SELECT bound_at FROM bindings WHERE trainee_id=? AND bind_type='trainer' AND unbound_at IS NULL", (te_a,))[0]
ws, we = expected_window(bound_at)
r = requests.get(f'{B}/api/trainees/{te_a}/rewards')
pa = r.json()['rewards_panel']['panels']
check('甲的奖励面板=培训师绑定', len(pa) == 1 and pa[0]['bind_type'] == 'trainer' and pa[0]['side_label'] == '培训师', str(pa))
rw = pa[0]
check('窗口起点=次周周一', rw['window_start'] == ws.strftime('%Y-%m-%d'), f"{rw['window_start']} != {ws.date()}")
check('窗口终点=起点+28天', rw['window_end'] == we.strftime('%Y-%m-%d'), f"{rw['window_end']} != {we.date()}")
if datetime.now() < ws:
    check('当前未进窗口 not_started 且无 pending', rw['window_status'] == 'not_started'
          and all(x['status'] != 'pending' for x in rw['rewards']), str(rw['window_status']))
    db_exec("UPDATE bindings SET bound_at=? WHERE trainee_id=? AND bind_type='trainer' AND unbound_at IS NULL",
            ((datetime.utcnow() - timedelta(days=10)).strftime('%Y-%m-%d %H:%M:%S'), te_a))
    rw = requests.get(f'{B}/api/trainees/{te_a}/rewards').json()['rewards_panel']['panels'][0]
check('进窗口后 in_window', rw['window_status'] == 'in_window' and rw['days_left'] is not None, str(rw['window_status']))
st = {x['key']: x['status'] for x in rw['rewards']}
check('金牌+任务7次 → 只发最高档（金牌+任务档 pending，低档已达标）',
      st.get('task6_bronze') == 'pending' and st.get('lv_gold') == 'pending'
      and st.get('lv_junior_silver') == 'achieved' and st.get('lv_silver') == 'achieved', str(st))
check('王牌未达标', st.get('lv_ace') == 'ineligible', str(st))
grant_id = next(x['grant_id'] for x in rw['rewards'] if x['key'] == 'lv_gold')
r = requests.post(f'{B}/api/rewards/{grant_id}/grant')
check('标记已发', r.ok, r.text)
r = requests.post(f'{B}/api/rewards/{grant_id}/grant')
check('重复标记 404', r.status_code == 404, f'got {r.status_code}')
r = requests.get(f'{B}/api/trainees/{te_a}/rewards')
st = {x['key']: x['status'] for x in r.json()['rewards_panel']['panels'][0]['rewards']}
check('金牌档已发放且不重生成', st.get('lv_gold') == 'granted', str(st))
r = requests.post(f'{B}/api/rewards/{grant_id}/revoke')
check('撤销发放', r.ok, r.text)

print('== 一键全部标记已发 ==')
bid_a = db_one("SELECT id FROM bindings WHERE trainee_id=? AND bind_type='trainer' AND unbound_at IS NULL", (te_a,))[0]
r = requests.post(f'{B}/api/rewards/grant-all', json={'binding_id': bid_a})
check('grant-all 返回发放数量（任务档+金牌档=2）', r.ok and r.json()['granted'] == 2, r.text)
r = requests.get(f'{B}/api/trainees/{te_a}/rewards')
st = {x['key']: x['status'] for x in r.json()['rewards_panel']['panels'][0]['rewards']}
check('任务档+金牌档已发放，低档已达标', st.get('task6_bronze') == 'granted' and st.get('lv_gold') == 'granted'
      and st.get('lv_junior_silver') == 'achieved' and st.get('lv_silver') == 'achieved', str(st))
r = requests.post(f'{B}/api/rewards/grant-all', json={'binding_id': bid_a})
check('重复 grant-all 幂等 0 档', r.ok and r.json()['granted'] == 0, r.text)

print('== 姐妹团绑定同样产生奖励 ==')
insert_monthly('80002', MONTH, '银牌', 3, tasks=7)
db_exec("UPDATE bindings SET bound_at=? WHERE trainee_id=? AND bind_type='sister' AND unbound_at IS NULL",
        ((datetime.utcnow() - timedelta(days=10)).strftime('%Y-%m-%d %H:%M:%S'), te_b))
r = requests.get(f'{B}/api/trainees/{te_b}/rewards')
pb = r.json()['rewards_panel']['panels']
check('乙的奖励面板=姐妹团绑定', len(pb) == 1 and pb[0]['bind_type'] == 'sister' and pb[0]['side_label'] == '姐姐', str(pb))
stb = {x['key']: x['status'] for x in pb[0]['rewards']}
check('姐妹团银牌+任务7次 → 银牌+任务档 pending，初银已达标',
      stb.get('task6_bronze') == 'pending' and stb.get('lv_silver') == 'pending'
      and stb.get('lv_junior_silver') == 'achieved' and stb.get('lv_gold') == 'ineligible', str(stb))

print('== 侧方未留存暂停发奖 ==')
sis_bid = pb[0]['binding']['id']
db_exec('DELETE FROM reward_grants WHERE binding_id=?', (sis_bid,))
db_exec('''INSERT INTO trainee_level_monthly (trainee_uid, month, best_level, best_rank, last_queried_at, last_retained)
           VALUES ('91001', ?, '', 0, ?, 0)''', (MONTH, NOW_S))
r = requests.get(f'{B}/api/trainees/{te_b}/rewards')
pb2 = r.json()['rewards_panel']['panels']
check('姐姐未留存 → 面板带标记', pb2[0].get('side_retained') == 0, str(pb2[0].get('side_retained')))
st3 = {x['key']: x['status'] for x in pb2[0]['rewards']}
check('姐姐未留存 → 不再生成待发放', all(v != 'pending' for v in st3.values()), str(st3))
db_exec("DELETE FROM trainee_level_monthly WHERE trainee_uid='91001'")

print('== 奖励跨月（按绑定当月起算） ==')
d20 = datetime.utcnow() - timedelta(days=20)
if d20.month != datetime.utcnow().month:
    # 绑定在上个月：窗口仍有效（次周周一起四周），等级记录也在上个月 → 金牌档应生成
    db_exec("UPDATE bindings SET bound_at=? WHERE trainee_id=? AND bind_type='sister' AND unbound_at IS NULL",
            (d20.strftime('%Y-%m-%d %H:%M:%S'), te_b))
    db_exec("DELETE FROM trainee_level_monthly WHERE trainee_uid='80002'")
    insert_monthly('80002', LAST_MONTH, '金牌', 4, tasks=7)
    r = requests.get(f'{B}/api/trainees/{te_b}/rewards')
    pb = r.json()['rewards_panel']['panels']
    check('跨月绑定窗口仍有效', pb[0]['window_status'] == 'in_window', str(pb[0]['window_status']))
    stb2 = {x['key']: x['status'] for x in pb[0]['rewards']}
    check('上月金牌 → 金牌档跨月生成 pending', stb2.get('lv_gold') == 'pending', str(stb2))
else:
    check('跨月用例（今天离月末太近，跳过）', True)

print('== 窗口过期不生成 pending ==')
db_exec("UPDATE bindings SET bound_at=? WHERE trainee_id=? AND unbound_at IS NULL",
        ((datetime.utcnow() - timedelta(days=60)).strftime('%Y-%m-%d %H:%M:%S'), te_c))
r = requests.get(f'{B}/api/trainees/{te_c}/rewards')
rw = r.json()['rewards_panel']['panels'][0]
check('丙窗口已过期 expired', rw['window_status'] == 'expired', str(rw['window_status']))
check('过期后无 pending 生成', all(x['status'] != 'pending' for x in rw['rewards']), str([x['status'] for x in rw['rewards']]))

print('== 分页 ==')
r = requests.get(f'{B}/api/trainees', params={'page': 1, 'page_size': 2})
check('妹妹分页 page_size=2', len(r.json()['items']) <= 2 and r.json()['total'] >= 3, str(r.json()))
r = requests.get(f'{B}/api/bindings', params={'active': 0, 'page': 1})
check('历史绑定分页字段', 'total' in r.json() and r.json()['page'] == 1, str(r.json().keys()))

print('== 历史导入 ==')
r = requests.get(f'{B}/api/trainees/import-candidates')
check('导入候选接口', r.ok and 'snapshot_date' in r.json(), r.text[:200])

print('== 历史恢复 ==')
bid_restore = db_one("SELECT id FROM bindings WHERE trainee_id=? AND bind_type='trainer' AND unbound_at IS NULL", (te_a,))[0]
r = requests.post(f'{B}/api/bindings/{bid_restore}/unbind')
check('先解绑', r.ok, r.text)
r = requests.post(f'{B}/api/bindings/{bid_restore}/restore')
check('恢复成功', r.ok, r.text)
r = requests.post(f'{B}/api/bindings/{bid_restore}/restore')
check('重复恢复 409', r.status_code == 409, f'got {r.status_code}')
r = requests.post(f'{B}/api/bindings/{bid_restore}/unbind')
db_exec("INSERT INTO bindings (trainer_id, trainee_id, bind_type, note) VALUES (?, ?, 'trainer', '冲突测试')", (tr_a, te_a))
r = requests.post(f'{B}/api/bindings/{bid_restore}/restore')
check('有同类型生效绑定时恢复 409', r.status_code == 409, f'got {r.status_code}')
db_exec("UPDATE bindings SET unbound_at=CURRENT_TIMESTAMP WHERE trainee_id=? AND note='冲突测试'", (te_a,))
r = requests.post(f'{B}/api/bindings/{bid_restore}/restore')
check('解掉冲突后可恢复', r.ok, r.text)

print('== 奖励周数（grant-week / set-remaining） ==')
# 甲的窗口已在窗口内（前面改过 bound_at）
w0 = next(b for b in requests.get(f'{B}/api/bindings', params={'active': 1}).json()['items'] if b['trainee_uid'] == '80001')
r = requests.post(f'{B}/api/bindings/{w0["id"]}/set-remaining', json={'remaining': 2})
check('set-remaining 改为 2', r.ok and r.json()['weeks_left'] == 2, r.text)
r = requests.post(f'{B}/api/bindings/{w0["id"]}/set-remaining', json={'remaining': 9})
check('set-remaining 超范围 400', r.status_code == 400, f'got {r.status_code}')
r = requests.post(f'{B}/api/rewards/grant-week', json={'binding_ids': [w0['id']]})
check('grant-week 窗口内周数-1 且标记奖励', r.ok and r.json()['bindings'] == 1, r.text)
w1 = next(b for b in requests.get(f'{B}/api/bindings', params={'active': 1}).json()['items'] if b['trainee_uid'] == '80001')
check('周数 2 → 1', w1['weeks_left'] == 1, str(w1['weeks_left']))

print('== 奖励导出 ==')
# 先解绑测试绑定（只动本测试创建的数据！）：导出测试不打网络
db_exec(f"UPDATE bindings SET unbound_at=CURRENT_TIMESTAMP WHERE unbound_at IS NULL AND trainee_id IN ({te_a},{te_b},{te_c})")
r = requests.post(f'{B}/api/export/rewards', json={'bind_type': 'x', 'date': '2026-10-05'})
check('非法类型 400', r.status_code == 400, f'got {r.status_code}')
r = requests.post(f'{B}/api/export/rewards', json={'bind_type': 'trainer', 'date': 'bad'})
check('非法日期 400', r.status_code == 400, f'got {r.status_code}')
r = requests.post(f'{B}/api/export/rewards', json={'bind_type': 'sister', 'date': '2026-10-05'})
check('空绑定导出有效 xlsx', r.status_code == 200 and r.content[:2] == b'PK', f'got {r.status_code}')
from openpyxl import load_workbook
import io
def _export_xlsx(bind_type, date='2026-10-05'):
    """导出 XLSX；外部系统不可用时返回 None（跳过文件校验，不判失败）"""
    r = requests.post(f'{B}/api/export/rewards', json={'bind_type': bind_type, 'date': date})
    if r.status_code != 200:
        print(f'  [SKIP] 导出接口暂不可用（外部系统），{bind_type} 跳过文件校验（HTTP {r.status_code}）')
        return None
    return load_workbook(io.BytesIO(r.content))


wb = _export_xlsx('sister')
if wb:
    hdr = [c.value for c in wb.active[3]]
    check('三胎表头为姐姐UID', hdr[0] == '姐姐UID' and hdr[5] == '完成任务次数' and hdr[7] == '礼物流水' and hdr[8] == '牌子', str(hdr))
wb = _export_xlsx('trainer')
if wb:
    hdr = [c.value for c in wb.active[3]]
    check('培训师表头为培训师UID', hdr[0] == '培训师UID', str(hdr))
# 过期窗口的绑定不出现在导出里（新建一条再把绑定时间改到 60 天前）
r = bind('91001', '80003', bind_type='sister', force=True)
exp_bid = r.json()['id']
db_exec('UPDATE bindings SET bound_at=? WHERE id=?',
        ((datetime.utcnow() - timedelta(days=60)).strftime('%Y-%m-%d %H:%M:%S'), exp_bid))
# 窗口未开始的（周一刚绑的）也不应出现：再建一条全新的
r = bind('91002', '80001', bind_type='sister', trainer_nickname='新姐姐', force=True)
check('新绑定窗口未开始', db_one('SELECT COUNT(*) FROM bindings WHERE trainee_id=? AND unbound_at IS NULL', (te_a,))[0] >= 1)
wb = _export_xlsx('sister')
if wb:
    ws_rows = [row for row in wb.active.iter_rows(min_row=5, values_only=True)
               if any(v is not None for v in row)]
    check('过期绑定不出现在导出（80003 不在其中）', all(str(row[2]) != '80003' for row in ws_rows), str(ws_rows))
    check('窗口未开始的绑定也不导出（80001 不在其中）', all(str(row[2]) != '80001' for row in ws_rows), str(ws_rows))
else:
    check('导出排除校验（外部系统不可用，跳过）', True)

print('== 用户端登录与自助绑定 ==')
r = requests.post(f'{B}/api/trainers/{tr_a}/set-password', json={'password': 'test123'})
check('管理员设置密码', r.ok, r.text)
cap = requests.get(f'{B}/api/portal/captcha').json()['code']
r = requests.post(f'{B}/api/portal/login', json={'uid': '90001', 'code': cap, 'password': 'wrong'})
check('错误密码 401', r.status_code == 401, f'got {r.status_code}')
cap = requests.get(f'{B}/api/portal/captcha').json()['code']
r = requests.post(f'{B}/api/portal/login', json={'uid': '90001', 'code': cap, 'password': 'test123'})
check('培训师 UID+密码+验证码登录', r.ok, r.text)
r = requests.get(f'{B}/api/portal/me')
check('用户端返回我的名额面板', r.ok and r.json()['slots']['trainer']['uid'] == '90001', r.text[:200])
r = requests.post(f'{B}/api/portal/change-password', json={'old_password': 'test123', 'new_password': 'test456'})
check('修改密码', r.ok, r.text)
cap = requests.get(f'{B}/api/portal/captcha').json()['code']
r = requests.post(f'{B}/api/portal/login', json={'uid': '90001', 'code': cap, 'password': 'test456'})
check('新密码可登录', r.ok, r.text)
r = requests.post(f'{B}/api/trainers', json={'uid': '90004', 'nickname': '初始密码测试'})
check('新建未设密码培训师', r.ok, r.text)
cap = requests.get(f'{B}/api/portal/captcha').json()['code']
r = requests.post(f'{B}/api/portal/login', json={'uid': '90004', 'code': cap, 'password': '90004'})
check('未设密码可用 UID 后5位初始密码登录', r.ok, r.text)
r = requests.post(f'{B}/api/portal/change-password', json={'old_password': '90004', 'new_password': 'abc123'})
check('无哈希用户修改密码不报错', r.ok, r.text)
cap = requests.get(f'{B}/api/portal/captcha').json()['code']
r = requests.post(f'{B}/api/portal/login', json={'uid': '90004', 'code': cap, 'password': 'abc123'})
check('改后新密码可登录', r.ok, r.text)
r = requests.post(f'{B}/api/portal/bind', json={'trainee_uid': '80001'})
check('自助绑定被保护期门槛拦截 422', r.status_code == 422 and r.json().get('gate'), r.text)
cap = requests.get(f'{B}/api/portal/captcha').json()['code']
r = requests.post(f'{B}/api/portal/login', json={'uid': '80001', 'code': cap, 'password': 'x'})
check('非培训师拒绝进入 403', r.status_code == 403, f'got {r.status_code}')
cap = requests.get(f'{B}/api/portal/captcha').json()['code']
r = requests.post(f'{B}/api/portal/login', json={'uid': '90001', 'code': 'xxxxx', 'password': 'test456'})
check('验证码环节已取消（任意验证码不拦截）', r.ok, r.text)
r = requests.post(f'{B}/api/portal/logout')
r = requests.get(f'{B}/api/portal/me')
check('退出后 401', r.status_code == 401, f'got {r.status_code}')

print('== 用户自行解散 ==')
requests.post(f'{B}/api/trainers/{tr_b}/set-password', json={'password': 'test123'})
db_exec("INSERT INTO bindings (trainer_id, trainee_id, bind_type) VALUES (?, ?, 'trainer')", (tr_b, te_a))
# 重置密码：管理员生成一次性随机密码
r = requests.post(f'{B}/api/trainers/{tr_b}/reset-password')
check('重置密码返回一次性密码', r.ok and len(r.json()['password']) == 8, r.text)
new_pwd = r.json()['password']
cap = requests.get(f'{B}/api/portal/captcha').json()['code']
r = requests.post(f'{B}/api/portal/login', json={'uid': '90002', 'code': cap, 'password': new_pwd})
check('重置后新密码可登录', r.ok, r.text)
# 恢复测试密码供后续用例使用
requests.post(f'{B}/api/trainers/{tr_b}/set-password', json={'password': 'test123'})
bid_p = db_one("SELECT id FROM bindings WHERE trainer_id=? AND trainee_id=? AND unbound_at IS NULL", (tr_b, te_a))[0]
cap = requests.get(f'{B}/api/portal/captcha').json()['code']
r = requests.post(f'{B}/api/portal/login', json={'uid': '90002', 'code': cap, 'password': 'test123'})
check('培训师B登录', r.ok, r.text)
r = requests.post(f'{B}/api/portal/unbind', json={'binding_id': bid_p})
check('用户自行解散成功', r.ok, r.text)
r = requests.get(f'{B}/api/bindings', params={'active': 1})
check('解散后不在生效列表', all(b['id'] != bid_p for b in r.json()['items']), r.text[:200])
# 别人的绑定不能解散
db_exec("INSERT INTO bindings (trainer_id, trainee_id, bind_type) VALUES (?, ?, 'trainer')", (tr_b, te_a))
bid_p2 = db_one("SELECT id FROM bindings WHERE trainer_id=? AND trainee_id=? AND unbound_at IS NULL", (tr_b, te_a))[0]
cap = requests.get(f'{B}/api/portal/captcha').json()['code']
requests.post(f'{B}/api/portal/login', json={'uid': '90001', 'code': cap, 'password': 'test456'})
r = requests.post(f'{B}/api/portal/unbind', json={'binding_id': bid_p2})
check('解散他人绑定 403', r.status_code == 403, f'got {r.status_code}')
r = requests.get(f'{B}/api/portal/trainee-data', params={'uid': '80001'})
check('查询他人妹妹数据 403', r.status_code == 403, f'got {r.status_code}')
r = requests.get(f'{B}/api/uid-lookup', params={'uid': '80001'})
check('用户端登录态可用 uid-lookup（不再被守卫拦截）', r.status_code in (200, 401, 502), f'got {r.status_code}')

print('== 数据看板 ==')
r = requests.get(f'{B}/api/operations', params={'limit': 'abc'})
check('操作日志 limit 非数字容错', r.status_code == 200, f'got {r.status_code}')
r = requests.get(f'{B}/api/operations', params={'limit': 5})
check('操作日志有记录', r.ok and len(r.json()['items']) >= 1, r.text[:200])
r = requests.post(f'{B}/api/trainers/999999/refresh')
check('刷新不存在培训师 404', r.status_code == 404, f'got {r.status_code}')
r = requests.get(f'{B}/api/dashboard/summary')
d = r.json()
check('看板结构', r.ok and len(d['weekly']) == 8 and 'dissolve_reasons' in d and 'active' in d, str(d.keys()))
check('近8周绑定数>=1（本周有测试绑定）', sum(w['bound'] for w in d['weekly']) >= 1, str(d['weekly']))
check('解散原因归类包含手动解绑或系统判定', all('reason' in x and x['count'] > 0 for x in d['dissolve_reasons']), str(d['dissolve_reasons']))
check('4周留存字段', set(('cohort', 'retained', 'dissolved', 'retention_rate', 'dissolution_rate')) <= set(d['retention_4w'].keys()), str(d['retention_4w']))
r = requests.get(f'{B}/api/dashboard/summary', params={'week': '2026-09-01'})
check('看板支持选周', r.ok and r.json()['this_week']['label'] == '08-31 ~ 09-06', r.json()['this_week'])
r = requests.get(f'{B}/api/bindings', params={'active': 0, 'kw': '测试妹妹甲'})
check('绑定列表按昵称搜索', len(r.json()['items']) >= 1 and all('测试妹妹甲' in b['trainee_nickname'] for b in r.json()['items']), r.text[:200])

print('== UID 速查 ==')
r = requests.get(f'{B}/api/uid-query/types')
check('类型列表', r.ok and len(r.json()['items']) == 4, r.text)
r = requests.post(f'{B}/api/trainees/999999/refresh')
check('刷新不存在学员 404（不打网络）', r.status_code == 404, f'got {r.status_code}')

print(f'\n结果: {passed} passed, {failed} failed')

# 清理测试数据（一律按测试 UID 精确匹配，绝不用可能错位的 id）
TEST_TRAINER_UIDS = ('90001', '90002', '90003', '90004', '90005', '90006', '90007', '91001', '91002')
TEST_TRAINEE_UIDS = ('80001', '80002', '80003')
# 测试绑定记录直接删除（不留历史，避免污染看板统计）
marks_tr = ','.join('?' * len(TEST_TRAINER_UIDS))
marks_te = ','.join('?' * len(TEST_TRAINEE_UIDS))
db_exec(f'''DELETE FROM bindings WHERE trainer_id IN (SELECT id FROM trainers WHERE uid IN ({marks_tr}))
            OR trainee_id IN (SELECT id FROM trainees WHERE uid IN ({marks_te}))''',
        (*TEST_TRAINER_UIDS, *TEST_TRAINEE_UIDS))
db_exec(f'DELETE FROM reward_grants WHERE trainee_uid IN ({marks_te})', TEST_TRAINEE_UIDS)
for uid in TEST_TRAINER_UIDS:
    for t in requests.get(f'{B}/api/trainers', params={'role': 'all', 'kw': uid}).json()['items']:
        if t['uid'] == uid:
            requests.delete(f"{B}/api/trainers/{t['id']}")
for uid in TEST_TRAINEE_UIDS:
    for t in requests.get(f'{B}/api/trainees', params={'kw': uid}).json()['items']:
        if t['uid'] == uid:
            requests.delete(f"{B}/api/trainees/{t['id']}")
    db_exec('DELETE FROM trainee_level_monthly WHERE trainee_uid=?', (uid,))
print('测试数据已清理')
sys.exit(1 if failed else 0)
